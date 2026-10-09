"""Document parsing and overlapping text chunking for the chatbot knowledge base."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath
from typing import Iterable

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".py",
    ".html",
    ".yaml",
    ".yml",
}


@dataclass(frozen=True)
class TextChunk:
    """A text fragment plus enough metadata to cite its origin."""

    text: str
    source: str
    chunk_id: int
    page: int | None = None


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 150) -> list[str]:
    """Split text into readable chunks while preserving a configurable overlap.

    Splits prefer a paragraph, sentence, or word boundary near the target size.
    The algorithm always advances, including when overlap is larger than a short
    final chunk, so malformed settings cannot create an infinite loop.
    """
    clean_text = "\n".join(line.rstrip() for line in (text or "").splitlines()).strip()
    if not clean_text:
        return []

    size = max(100, int(chunk_size))
    overlap_size = min(max(0, int(overlap)), size - 1)
    chunks: list[str] = []
    start = 0
    text_length = len(clean_text)

    while start < text_length:
        target_end = min(start + size, text_length)
        end = target_end

        if target_end < text_length:
            search_floor = start + max(1, size // 2)
            boundaries = [
                clean_text.rfind("\n", search_floor, target_end),
                clean_text.rfind(". ", search_floor, target_end),
                clean_text.rfind("? ", search_floor, target_end),
                clean_text.rfind("! ", search_floor, target_end),
                clean_text.rfind(" ", search_floor, target_end),
            ]
            boundary = max(boundaries)
            if boundary >= search_floor:
                end = boundary + (1 if clean_text[boundary : boundary + 1] == "\n" else 0)
                end = max(start + 1, end)

        piece = clean_text[start:end].strip()
        if piece:
            chunks.append(piece)

        if end >= text_length:
            break

        next_start = end - overlap_size
        if next_start <= start:
            next_start = min(text_length, start + size)
        start = next_start

    return chunks


def _extract_pdf(content: bytes) -> list[tuple[str, int | None]]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        return [(page.extract_text() or "", page_number) for page_number, page in enumerate(reader.pages, start=1)]
    except Exception as exc:  # parser errors vary by PDF type and pypdf version
        raise ValueError(f"Could not read this PDF: {exc}") from exc


def _extract_docx(content: bytes) -> list[tuple[str, int | None]]:
    try:
        from docx import Document

        document = Document(BytesIO(content))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        table_rows: list[str] = []
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(values):
                    table_rows.append(" | ".join(values))
        combined = "\n".join(paragraphs + table_rows)
        return [(combined, None)]
    except Exception as exc:
        raise ValueError(f"Could not read this DOCX file: {exc}") from exc


def extract_text_pages(filename: str, content: bytes) -> list[tuple[str, int | None]]:
    """Extract text from a supported upload; PDF pages retain page numbers."""
    suffix = PurePath(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type '{suffix or '(no extension)'}'. Supported types: {supported}")
    if not content:
        raise ValueError("The uploaded file is empty.")

    if suffix == ".pdf":
        pages = _extract_pdf(content)
    elif suffix == ".docx":
        pages = _extract_docx(content)
    else:
        try:
            decoded = content.decode("utf-8-sig", errors="replace")
        except Exception as exc:
            raise ValueError(f"Could not decode this file: {exc}") from exc
        pages = [(decoded, None)]

    if not any(text.strip() for text, _ in pages):
        raise ValueError("No extractable text was found. For scanned PDFs, run OCR before uploading.")
    return pages


def parse_file_to_chunks(
    filename: str,
    content: bytes,
    chunk_size: int = 1000,
    overlap: int = 150,
) -> list[TextChunk]:
    """Parse a user file and split it into citable chunks."""
    pages = extract_text_pages(filename, content)
    result: list[TextChunk] = []
    chunk_number = 1
    for page_text, page_number in pages:
        for piece in chunk_text(page_text, chunk_size=chunk_size, overlap=overlap):
            result.append(
                TextChunk(
                    text=piece,
                    source=PurePath(filename).name,
                    chunk_id=chunk_number,
                    page=page_number,
                )
            )
            chunk_number += 1
    if not result:
        raise ValueError("No searchable text chunks could be created from this file.")
    return result


def chunks_to_plain_text(chunks: Iterable[TextChunk]) -> str:
    """Small utility useful for debugging and testing parsed documents."""
    return "\n\n".join(chunk.text for chunk in chunks)
